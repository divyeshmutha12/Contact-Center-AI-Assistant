"""
Test script for MariaDB Query Tool

Run this script independently to verify:
1. MariaDB connection works
2. Queries execute correctly
3. Excel export works

Usage:
    python test_mariadb_tool.py
    python test_mariadb_tool.py --query agent_status_report
    python test_mariadb_tool.py --list
"""

import os
import sys
import json
import argparse
from datetime import datetime
from dotenv import load_dotenv

# Load environment variables
load_dotenv()

# Check for required packages
try:
    import pymysql
    from pymysql.cursors import DictCursor
except ImportError:
    print("ERROR: pymysql not installed. Run: pip install pymysql")
    sys.exit(1)

try:
    from openpyxl import Workbook
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip install openpyxl")
    sys.exit(1)


def get_credentials():
    """Get MariaDB credentials from .env"""
    host = os.getenv("MYSQL_HOST")
    user = os.getenv("MYSQL_USER")
    password = os.getenv("MYSQL_PASS")
    database = os.getenv("MYSQL_DB_NAME", "ccs_dev")

    print("\n=== MariaDB Credentials ===")
    print(f"Host: {host}")
    print(f"User: {user}")
    print(f"Password: {'*' * len(password) if password else 'NOT SET'}")
    print(f"Database: {database}")

    if not all([host, user, password]):
        print("\nERROR: Missing credentials in .env file!")
        print("Required: MYSQL_HOST, MYSQL_USER, MYSQL_PASS")
        return None

    return {
        "host": host,
        "user": user,
        "password": password,
        "database": database
    }


def test_connection(creds):
    """Test basic MariaDB connection"""
    print("\n=== Testing Connection ===")
    try:
        connection = pymysql.connect(
            host=creds["host"],
            user=creds["user"],
            password=creds["password"],
            database=creds["database"],
            cursorclass=DictCursor,
            connect_timeout=10,
            read_timeout=300,
            charset='utf8mb4'
        )

        with connection.cursor() as cursor:
            cursor.execute("SELECT 1 as test")
            result = cursor.fetchone()
            print(f"Connection test: {result}")

            cursor.execute("SELECT VERSION() as version")
            version = cursor.fetchone()
            print(f"MariaDB Version: {version['version']}")

            cursor.execute("SELECT DATABASE() as db")
            db = cursor.fetchone()
            print(f"Current Database: {db['db']}")

        print("Connection: SUCCESS")
        return connection

    except pymysql.Error as e:
        print(f"Connection: FAILED")
        print(f"Error: {e}")
        return None


def load_queries_config():
    """Load queries from config file"""
    config_path = "config/mariadb_queries.json"
    if not os.path.exists(config_path):
        print(f"ERROR: Config file not found: {config_path}")
        return None

    with open(config_path, "r") as f:
        return json.load(f)


def list_queries(queries):
    """List all available queries"""
    print("\n=== Available Queries ===")
    for name, config in queries.items():
        desc = config.get("description", "No description")
        params = list(config.get("parameters", {}).keys())
        print(f"\n- {name}")
        print(f"  Description: {desc}")
        print(f"  Parameters: {params}")


def get_current_date_range():
    """Get current date start and end times"""
    today = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
    start_date = today.strftime("%Y-%m-%d 00:00:00")
    end_date = today.strftime("%Y-%m-%d 23:59:59")
    return start_date, end_date


def apply_current_date_defaults(parameters):
    """Replace CURRENT_DATE placeholders with actual current date"""
    current_start, current_end = get_current_date_range()

    for key, value in parameters.items():
        if value == "{{CURRENT_DATE_START}}":
            parameters[key] = current_start
        elif value == "{{CURRENT_DATE_END}}":
            parameters[key] = current_end

    return parameters


def replace_placeholders(query, parameters):
    """Replace placeholders in SQL query with actual values"""
    result = query
    for param_name, value in parameters.items():
        placeholder = "{{" + param_name + "}}"
        if placeholder in result:
            if isinstance(value, str):
                safe_value = value.replace("'", "''")
                result = result.replace(placeholder, safe_value)
            else:
                result = result.replace(placeholder, str(value))
    return result


def format_cell_value(value):
    """Format a value for Excel cell"""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d %H:%M:%S')
    if isinstance(value, bytes):
        return value.decode('utf-8', errors='ignore')
    if isinstance(value, (int, float)) and abs(value) > 1000000:
        return str(int(value))
    return value


def execute_query(connection, query_name, queries, param_overrides=None):
    """Execute a query from config"""
    print(f"\n=== Executing Query: {query_name} ===")

    if query_name not in queries:
        print(f"ERROR: Query '{query_name}' not found!")
        list_queries(queries)
        return None

    query_config = queries[query_name]

    # Build parameters
    parameters = query_config.get("parameters", {}).copy()

    # Apply current date defaults if configured
    if query_config.get("use_current_date_default", False):
        parameters = apply_current_date_defaults(parameters)

    # Apply any overrides
    if param_overrides:
        parameters.update(param_overrides)

    print(f"Parameters: {json.dumps(parameters, indent=2)}")

    # Build query
    query_template = query_config.get("query", "")
    query = replace_placeholders(query_template, parameters)

    # Print query (truncated for readability)
    print(f"\nSQL Query (first 500 chars):")
    print(query[:500] + "..." if len(query) > 500 else query)

    # Execute
    try:
        with connection.cursor() as cursor:
            print("\nExecuting query...")
            cursor.execute(query)
            results = cursor.fetchall()
            print(f"Query returned: {len(results)} rows")
            return results
    except pymysql.Error as e:
        print(f"Query execution FAILED: {e}")
        return None


def export_to_excel(results, output_filename):
    """Export results to Excel"""
    if not results:
        print("No results to export")
        return None

    print(f"\n=== Exporting to Excel: {output_filename} ===")

    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title="Report")

    # Get headers
    headers = list(results[0].keys())
    ws.append(headers)
    print(f"Headers: {headers}")

    # Write data
    rows_written = 0
    for row in results:
        row_data = [format_cell_value(row.get(key)) for key in headers]
        ws.append(row_data)
        rows_written += 1

    wb.save(output_filename)
    wb.close()

    print(f"Excel saved: {output_filename} ({rows_written} rows)")
    return output_filename


def print_sample_results(results, num_rows=5):
    """Print sample results"""
    if not results:
        return

    print(f"\n=== Sample Results (first {num_rows} rows) ===")
    for i, row in enumerate(results[:num_rows]):
        print(f"\n--- Row {i+1} ---")
        for key, value in row.items():
            formatted_value = format_cell_value(value)
            # Truncate long values
            if isinstance(formatted_value, str) and len(formatted_value) > 50:
                formatted_value = formatted_value[:50] + "..."
            print(f"  {key}: {formatted_value}")


def main():
    parser = argparse.ArgumentParser(description="Test MariaDB Query Tool")
    parser.add_argument("--query", "-q", help="Query name to execute")
    parser.add_argument("--list", "-l", action="store_true", help="List available queries")
    parser.add_argument("--start-date", help="Override start date (YYYY-MM-DD HH:MM:SS)")
    parser.add_argument("--end-date", help="Override end date (YYYY-MM-DD HH:MM:SS)")
    parser.add_argument("--sme-id", type=int, help="Override SME ID")
    parser.add_argument("--output", "-o", help="Output Excel filename")
    parser.add_argument("--no-export", action="store_true", help="Skip Excel export")

    args = parser.parse_args()

    print("=" * 60)
    print("MariaDB Query Tool - Test Script")
    print("=" * 60)

    # Get credentials
    creds = get_credentials()
    if not creds:
        sys.exit(1)

    # Test connection
    connection = test_connection(creds)
    if not connection:
        sys.exit(1)

    # Load queries config
    queries = load_queries_config()
    if not queries:
        connection.close()
        sys.exit(1)

    # List queries if requested
    if args.list:
        list_queries(queries)
        connection.close()
        sys.exit(0)

    # Execute query if specified
    if args.query:
        # Build parameter overrides
        param_overrides = {}
        if args.start_date:
            param_overrides["start_date"] = args.start_date
        if args.end_date:
            param_overrides["end_date"] = args.end_date
        if args.sme_id:
            param_overrides["sme_id"] = args.sme_id

        results = execute_query(connection, args.query, queries, param_overrides)

        if results:
            print_sample_results(results)

            if not args.no_export:
                output_file = args.output or f"test_{args.query}_report.xlsx"
                export_to_excel(results, output_file)
    else:
        # Default: just test connection and list queries
        list_queries(queries)
        print("\n" + "=" * 60)
        print("To execute a query, use: python test_mariadb_tool.py --query <query_name>")
        print("Example: python test_mariadb_tool.py --query sms_report")
        print("=" * 60)

    # Cleanup
    connection.close()
    print("\nConnection closed.")
    print("Test completed.")


if __name__ == "__main__":
    main()
